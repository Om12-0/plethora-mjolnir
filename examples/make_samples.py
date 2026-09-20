"""Generate sample .pdf and .docx files for trying PLETHORA MJOLNIR.

Run:  python examples/make_samples.py
Creates: examples/sample_docs/circuit_report.pdf  and  assignment3.docx
Pure standard library for the PDF (no reportlab needed).
"""
import os

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "sample_docs")


def _esc(s):
    return s.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def build_pdf(path, pages):
    """Minimal valid multi-page PDF built by hand (Helvetica, text only)."""
    objs = []

    def add(body):
        objs.append(body)
        return len(objs)

    catalog = add(b"")                                    # 1
    pages_obj = add(b"")                                  # 2
    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    page_nums, content_nums = [], []
    for _ in pages:
        content_nums.append(add(b""))
        page_nums.append(add(b""))

    objs[pages_obj - 1] = (b"<< /Type /Pages /Kids ["
                           + b" ".join(b"%d 0 R" % p for p in page_nums)
                           + b"] /Count %d >>" % len(page_nums))
    objs[catalog - 1] = b"<< /Type /Catalog /Pages %d 0 R >>" % pages_obj

    for idx, lines in enumerate(pages):
        ops = ["BT", "/F1 12 Tf", "14 TL", "72 720 Td"]
        for i, line in enumerate(lines):
            if i:
                ops.append("T*")
            ops.append("(%s) Tj" % _esc(line))
        ops.append("ET")
        content = ("\n".join(ops)).encode("latin-1", "replace")
        objs[content_nums[idx] - 1] = (b"<< /Length %d >>\nstream\n" % len(content)
                                       + content + b"\nendstream")
        objs[page_nums[idx] - 1] = (
            b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] "
            b"/Contents %d 0 R /Resources << /Font << /F1 %d 0 R >> >> >>"
            % (pages_obj, content_nums[idx], font)
        )

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n" % (len(objs) + 1) + b"0000000000 65535 f \n"
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += (b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n"
            % (len(objs) + 1, catalog, xref))
    with open(path, "wb") as fh:
        fh.write(bytes(out))


def build_docx(path):
    import docx
    d = docx.Document()
    d.add_heading("Assignment 3 - Database Normalization", level=1)
    d.add_paragraph(
        "Task: design a restaurant ordering database schema and justify it "
        "up to third normal form (3NF)."
    )
    d.add_heading("Deliverables", level=2)
    for line in [
        "Entity relationship description",
        "CREATE TABLE statements (see restaurant_schema.sql)",
        "Discussion of partial and transitive dependencies",
    ]:
        d.add_paragraph(line, style="List Bullet")
    d.add_heading("Grading notes", level=2)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Criteria"
    t.cell(0, 1).text = "Weight"
    t.cell(1, 0).text = "Correct normalization"
    t.cell(1, 1).text = "60%"
    d.add_paragraph(
        "Reminder: the June capacitor lab report is a separate submission and "
        "is not part of this assignment."
    )
    d.save(path)


def main():
    os.makedirs(OUT, exist_ok=True)
    pdf_path = os.path.join(OUT, "circuit_report.pdf")
    build_pdf(pdf_path, pages=[
        [
            "PHY2049 - Lab 3: RC Circuit Capacitor Discharge",
            "Date: June 14, 2025",
            "",
            "Objective: measure the time constant of an RC circuit by",
            "recording capacitor voltage during discharge and fitting an",
            "exponential decay V(t) = V0 * exp(-t / RC).",
        ],
        [
            "Results",
            "",
            "Measured tau = 0.98 s for R = 10 kOhm and C = 100 uF.",
            "The exponential fit matches the data to within 2%.",
            "",
            "Conclusion: the capacitor discharge follows the expected",
            "exponential law and the time constant agrees with theory.",
        ],
    ])
    docx_path = os.path.join(OUT, "assignment3.docx")
    build_docx(docx_path)
    print("Wrote:")
    print("  " + pdf_path)
    print("  " + docx_path)


if __name__ == "__main__":
    main()
